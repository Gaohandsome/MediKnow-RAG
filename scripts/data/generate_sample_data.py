"""生成多格式解析测试样本（Markdown / PDF / PPTX / 图片）。

注意：演示主文档已改为《医院员工培训守则》PDF，由 generate_training_handbook.py 生成；
本脚本仅保留用于生成多格式（md / pdf / pptx / png）的解析回归样本，不参与默认演示流程。

说明：全部内容为**合成演示数据**，结构参考真实院内资料，但不具备临床效力，
不可用于实际诊疗决策。用法：python scripts/data/generate_sample_data.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "data"

DISCLAIMER_LINE = "【合成演示数据】本文件由项目脚本生成，仅用于系统演示与测试，不可用于临床决策。"

PNEUMONIA_MD = f"""# 社区获得性肺炎诊疗规范（院内试行版）

> {DISCLAIMER_LINE}

## 一、诊断依据

1. 临床表现：新近出现的咳嗽、咳痰，或原有呼吸道症状加重，伴或不伴胸痛、呼吸困难、发热。
2. 体征：肺实变体征或闻及湿性啰音。
3. 影像学：胸部 X 线或 CT 显示新出现的斑片状浸润影、叶段实变或间质性改变。
4. 实验室：白细胞计数升高或降低，降钙素原（PCT）可辅助判断细菌感染可能性。

具备第 3 条影像学依据，同时满足第 1 条或第 2 条中任一项，并排除肺结核、肺栓塞、肺水肿等疾病后，可临床诊断。

## 二、严重程度评估

采用 CURB-65 评分，每项 1 分：

- C：意识障碍
- U：血尿素氮 > 7 mmol/L
- R：呼吸频率 ≥ 30 次/分
- B：收缩压 < 90 mmHg 或舒张压 ≤ 60 mmHg
- 65：年龄 ≥ 65 岁

评分 0~1 分：可门诊治疗；2 分：建议住院治疗；≥ 3 分：评估入住重症监护病房。

## 三、经验性抗菌治疗

1. 门诊、无基础疾病患者：阿莫西林 500 mg 口服，每 8 小时一次。
2. 门诊、有慢性心肺疾病或近期使用抗菌药物史：阿莫西林/克拉维酸 875/125 mg 口服，每 12 小时一次。
3. 住院非重症患者：头孢曲松 1 g 静脉滴注，每日一次，联合阿奇霉素 0.5 g 静脉滴注，每日一次。
4. 重症患者：在住院方案基础上，按院内会诊意见调整，并由副主任医师及以上职称医师开具。

用药前需确认过敏史；青霉素类过敏者改用呼吸喹诺酮类，具体品种与剂量以药品说明书及院内抗菌药物目录为准。

## 四、疗程与停药指征

- 轻中度患者疗程一般为 5~7 天。
- 体温正常超过 48~72 小时，且呼吸、心率、血压等临床指标稳定，可考虑停药。
- 重症或存在肺外并发症者，疗程由主管医师根据复查结果个体化确定。

## 五、随访与复查

1. 出院后 7~10 天至呼吸内科门诊复诊，评估症状与体征。
2. 影像学吸收较慢，建议在 4~6 周后复查胸部影像，用于排除肿瘤等基础病变。
3. 吸烟患者需同时接受戒烟干预，并建议接种肺炎链球菌疫苗与流感疫苗。
"""

ANTIBIOTIC_CONTENT = [
    ("第一章 总则", "h"),
    ("第一条 为规范院内抗菌药物临床应用，遏制细菌耐药，依据《抗菌药物临床应用管理办法》制定本制度。", "p"),
    ("第二条 本制度适用于全院临床科室、药学部与医院感染管理部门。", "p"),
    ("第三条 抗菌药物临床应用实行分级管理，分为非限制使用级、限制使用级与特殊使用级。", "p"),
    ("第二章 分级原则", "h"),
    ("第四条 非限制使用级：疗效确切、安全性好、价格相对低廉，对细菌耐药性影响较小。", "p"),
    ("第五条 限制使用级：疗效确切但价格较高，或对细菌耐药性影响较大，或安全性需进一步观察。", "p"),
    ("第六条 特殊使用级：具有明显抗菌优势的新药，价格昂贵，或安全性证据尚不充分。", "p"),
    ("第三章 处方权限", "h"),
    ("第七条 非限制使用级：取得处方权的医师均可开具。", "p"),
    ("第八条 限制使用级：由主治医师及以上职称医师开具。", "p"),
    ("第九条 特殊使用级：需经院内抗菌药物会诊专家组会诊同意，由具有高级职称的医师开具。", "p"),
    ("第十条 紧急情况下越级使用特殊使用级药物，仅限一日用量，并于 24 小时内补办会诊与审批手续。", "p"),
    ("第四章 会诊与审批流程", "h"),
    ("第十一条 申请科室在 HIS 系统提交会诊申请，注明感染部位、病原学结果与用药理由。", "p"),
    ("第十二条 会诊专家组由感染科、呼吸科、重症医学科、临床药学与微生物室人员组成。", "p"),
    ("第十三条 会诊意见应在 24 小时内返回，紧急会诊应在 2 小时内完成。", "p"),
    ("第五章 监督与考核", "h"),
    ("第十四条 药学部每月统计各科室抗菌药物使用强度，并在院内质量简报中公示。", "p"),
    ("第十五条 越级使用未补办手续、无指征联合用药的，纳入科室质控扣分并全院通报。", "p"),
    ("第十六条 围手术期预防用药原则上不超过 24 小时，心脏与关节置换手术可延长至 48 小时。", "p"),
]

INFECTION_SLIDES = [
    ("院感防控培训课件（2026 年度）", [DISCLAIMER_LINE, "医院感染管理科 · 全员必修"]),
    (
        "一、手卫生五个时刻",
        [
            "接触患者前",
            "进行清洁或无菌操作前",
            "暴露于体液风险后",
            "接触患者后",
            "接触患者周围环境后",
        ],
    ),
    (
        "二、手卫生方法",
        [
            "洗手：流动水与皂液，全程不少于 40 秒，七步洗手法",
            "速干手消毒剂：无可见污染物时首选，全程 20~30 秒",
            "戴手套不能替代手卫生，脱手套后仍需洗手",
        ],
    ),
    (
        "三、接触隔离要求",
        [
            "多重耐药菌患者单间安置或同类集中安置",
            "进入病室须穿隔离衣、戴手套，出室前脱卸并洗手",
            "专用听诊器、体温计，物品表面每日含氯消毒剂擦拭",
            "患者转科或出院后执行终末消毒",
        ],
    ),
    (
        "四、职业暴露处置",
        [
            "锐器伤：一挤、二冲、三消毒、四上报",
            "立即从近心端向远心端挤出血液，禁止局部按压伤口",
            "流动水冲洗 5 分钟，75% 乙醇或碘伏消毒",
            "2 小时内上报院感科，评估暴露源并启动预防用药",
        ],
    ),
]


def _font_path() -> str | None:
    """返回第一个存在的中文字体路径（跨平台），找不到返回 None。"""
    for candidate in (
        os.environ.get("MEDIKNOW_FONT", ""),
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Songti.ttc",
        "/Library/Fonts/Arial Unicode.ttf",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
    ):
        if candidate and Path(candidate).exists():
            return candidate
    return None


def write_markdown() -> Path:
    path = DATA_DIR / "社区获得性肺炎诊疗规范.md"
    path.write_text(PNEUMONIA_MD, encoding="utf-8")
    return path


def write_pdf() -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    title_style = ParagraphStyle("title", fontName="STSong-Light", fontSize=18, leading=26, spaceAfter=10)
    note_style = ParagraphStyle("note", fontName="STSong-Light", fontSize=10, leading=16, textColor="#666666")
    head_style = ParagraphStyle("head", fontName="STSong-Light", fontSize=13, leading=22, spaceBefore=10)
    body_style = ParagraphStyle("body", fontName="STSong-Light", fontSize=11, leading=20)

    path = DATA_DIR / "抗菌药物分级管理制度.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4, title="抗菌药物分级管理制度")
    story = [
        Paragraph("院内抗菌药物分级管理制度", title_style),
        Paragraph(DISCLAIMER_LINE, note_style),
        Spacer(1, 6),
    ]
    for text, style_key in ANTIBIOTIC_CONTENT:
        style = head_style if style_key == "h" else body_style
        if style_key == "h":
            story.append(Spacer(1, 6))
        story.append(Paragraph(text, style))
    doc.build(story)
    return path


def write_pptx() -> Path:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    for title, bullets in INFECTION_SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        text_frame = slide.placeholders[1].text_frame
        text_frame.text = bullets[0]
        for bullet in bullets[1:]:
            paragraph = text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(18)
        slide.shapes.title.text_frame.paragraphs[0].font.size = Pt(30)
    path = DATA_DIR / "院感防控培训课件.pptx"
    prs.save(str(path))
    return path


def write_image() -> Path:
    from PIL import Image, ImageDraw, ImageFont

    width, height = 980, 620
    image = Image.new("RGB", (width, height), (246, 250, 252))
    draw = ImageDraw.Draw(image)

    font_path = _font_path()
    title_font = ImageFont.truetype(font_path, 32) if font_path else ImageFont.load_default()
    body_font = ImageFont.truetype(font_path, 19) if font_path else ImageFont.load_default()
    small_font = ImageFont.truetype(font_path, 15) if font_path else ImageFont.load_default()

    draw.rounded_rectangle(
        [36, 36, width - 36, height - 36], radius=22, fill=(255, 255, 255), outline=(214, 226, 234), width=2
    )
    draw.text((72, 66), "门诊就诊与检查流程", font=title_font, fill=(24, 48, 60))
    draw.text((72, 112), DISCLAIMER_LINE, font=small_font, fill=(130, 138, 146))

    steps = [
        ("1. 挂号", "自助机 / 公众号 / 窗口"),
        ("2. 分诊", "测体温、血压，分级候诊"),
        ("3. 就诊", "出示医保卡与既往病历"),
        ("4. 检查", "持申请单至相应科室"),
        ("5. 取报告", "自助机打印或线上查询"),
        ("6. 复诊", "携报告回原诊室"),
    ]
    box_w, box_h, gap = 128, 96, 18
    start_x, start_y = 72, 190
    for index, (title, subtitle) in enumerate(steps):
        x = start_x + index * (box_w + gap)
        draw.rounded_rectangle([x, start_y, x + box_w, start_y + box_h], radius=12, fill=(232, 244, 249), outline=(126, 178, 198), width=1)
        draw.text((x + 14, start_y + 16), title, font=body_font, fill=(20, 62, 78))
        draw.text((x + 14, start_y + 48), subtitle, font=small_font, fill=(76, 104, 116))
        if index < len(steps) - 1:
            draw.line([x + box_w + 2, start_y + box_h // 2, x + box_w + gap - 4, start_y + box_h // 2], fill=(126, 178, 198), width=2)

    notes = [
        "采血与腹部超声需空腹 8 小时以上，可少量饮水；糖尿病患者请随身携带糖果。",
        "增强 CT 检查前需确认肾功能与过敏史，检查后留观 30 分钟。",
        "报告出具时间：血常规约 30 分钟，生化约 2 小时，病理约 3 个工作日。",
        "危急值由检验科电话通知临床科室，患者请保持预留电话畅通。",
    ]
    y = 330
    for note in notes:
        draw.text((72, y), "·  " + note, font=body_font, fill=(60, 74, 84))
        y += 42

    path = DATA_DIR / "门诊就诊与检查流程图.png"
    image.save(path)
    return path


def main() -> int:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    created = [write_markdown(), write_pdf(), write_pptx(), write_image()]
    for path in created:
        print(f"已生成：{path.name}（{path.stat().st_size / 1024:.1f} KB）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
