"""文档解析：把各类离线文件转换成带页码的文本块，并顺带裁出版面里的图/表。"""
from __future__ import annotations

import base64
import re
from dataclasses import dataclass, field
from pathlib import Path

from loguru import logger

from app.core.config import settings
from app.core.exceptions import BizError

SUPPORTED_EXTENSIONS = {
    ".pdf": "PDF",
    ".pptx": "PPT",
    ".ppt": "PPT",
    ".docx": "Word",
    ".md": "Markdown",
    ".markdown": "Markdown",
    ".txt": "文本",
    ".png": "图片",
    ".jpg": "图片",
    ".jpeg": "图片",
    ".bmp": "图片",
    ".webp": "图片",
}

# 图注行：图 1 / 图1-2 / 图 3 流程图
_CAPTION_LINE = re.compile(r"^图\s*\d+")

# 一页抽出的文字少于该长度，且页面上有整页位图 → 判为「无文字层」（扫描/栅格化页面）
_SCAN_TEXT_MIN = 12

# 扫描页转录提示词：要求保留结构，而不是写摘要
_SCAN_PROMPT = (
    "这是医院制度文档中某一页的图像。请把页面内容完整转录成纯文本，要求：\n"
    "1. 标题行以 # / ## 开头，保持原有层级；\n"
    "2. 表格必须输出为 Markdown 表格，保留全部行列与单元格文字，不要省略、不要合并；\n"
    "3. 流程图/示意图按「步骤 1 → 步骤 2 …」逐框列出框内文字，并说明箭头指向；\n"
    "4. 只输出转录结果，不要总结、不要解释、不要写「以下是」。"
)


@dataclass
class TextBlock:
    """解析出的文本块，page 用于保留原始页/页签信息。

    images 记录该块关联的图片：`{"caption", "url", "kind"}`，kind 为 table/figure/image。
    """

    text: str
    page: int | None = None
    meta: dict = field(default_factory=dict)
    images: list[dict] = field(default_factory=list)


def get_file_type(filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    return SUPPORTED_EXTENSIONS.get(suffix, "")


def is_supported(filename: str) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def parse_file(
    path: str | Path, filename: str | None = None, document_id: int | None = None
) -> list[TextBlock]:
    """按文件类型分发解析器。

    document_id 用于给裁出的图片命名（`{document_id}/p4_fig1.png`）。
    """
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if not file_path.exists():
        raise BizError(f"文件不存在：{file_path}", status_code=404)

    if suffix == ".pdf":
        return _parse_pdf(file_path, document_id)
    if suffix == ".pptx":
        return _parse_pptx(file_path)
    if suffix == ".ppt":
        raise BizError("暂不支持 .ppt 旧格式，请另存为 .pptx 后重新上传")
    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix in {".md", ".markdown", ".txt"}:
        return _parse_text(file_path)
    if suffix in {".png", ".jpg", ".jpeg", ".bmp", ".webp"}:
        return _parse_image(file_path)
    raise BizError(f"不支持的文件类型：{suffix}")


def _table_to_markdown(rows: list[list[str | None]]) -> str:
    """把 pdfplumber 抽出的二维单元格转成 Markdown 表格，保留行列结构。"""
    cleaned: list[list[str]] = []
    for row in rows or []:
        cells = [(cell or "").replace("\n", " ").strip() for cell in row]
        if any(cells):
            cleaned.append(cells)
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    cleaned = [row + [""] * (width - len(row)) for row in cleaned]
    header, body = cleaned[0], cleaned[1:]
    lines = ["| " + " | ".join(header) + " |", "| " + " | ".join(["---"] * width) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _overlap_ratio(a: tuple, b: tuple) -> float:
    """两个 bbox 的重叠面积占较小者的比例。"""
    x0, top = max(a[0], b[0]), max(a[1], b[1])
    x1, bottom = min(a[2], b[2]), min(a[3], b[3])
    if x1 <= x0 or bottom <= top:
        return 0.0
    inter = (x1 - x0) * (bottom - top)
    area_a = max((a[2] - a[0]) * (a[3] - a[1]), 1e-6)
    area_b = max((b[2] - b[0]) * (b[3] - b[1]), 1e-6)
    return inter / min(area_a, area_b)


def _figure_bbox(page, caption: dict) -> tuple | None:
    """从图注向上聚类矢量对象，得到流程图/示意图的区域。

    步骤：先取紧贴图注上方的对象，再沿垂直方向按「间隙 <= 40pt」逐个向上吸收，
    直到断开为止——这样一个由多个框和箭头组成的流程图会被整体框住。
    """
    drawings = []
    for attr in ("rects", "curves", "lines"):
        drawings.extend(getattr(page, attr, []) or [])
    drawings = [item for item in drawings if item.get("bottom", 0) <= caption["top"] + 1]
    if not drawings:
        return None

    seed = max(drawings, key=lambda item: item["bottom"])
    if caption["top"] - seed["bottom"] > 60:
        return None

    group = [seed]
    limit = seed["top"]
    while True:
        candidates = [
            item
            for item in drawings
            if item not in group and item["bottom"] >= limit - 40
        ]
        if not candidates:
            break
        group.extend(candidates)
        limit = min(item["top"] for item in group)

    top = min(item["top"] for item in group)
    x0 = min(item["x0"] for item in group)
    x1 = max(item["x1"] for item in group)
    # 把图注一起裁进去，模型能读到「图 1 xxx」知道这张图是什么
    return (x0 - 6, max(top - 6, 0), x1 + 6, min(caption["bottom"] + 4, page.height))


def _crop_and_save(page, bbox: tuple, key: str) -> str | None:
    """按 bbox 裁图并保存，返回可访问 URL。"""
    try:
        image = page.crop(bbox).to_image(resolution=settings.media_crop_resolution).original
        from app.services.media import save_image

        return save_image(image, key)
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"裁图失败：{key} - {exc}")
        return None


def _transcribe_image(image, hint: str = "") -> str:
    """调用多模态模型把图像转录成文本（表格转 Markdown、流程图按步骤描述）。"""
    if not settings.vision_enabled:
        return ""
    from app.services.llm import complete
    from app.services.media import image_data_url

    prompt = _SCAN_PROMPT + (f"\n补充线索：{hint}" if hint else "")
    try:
        return complete(
            messages=[
                {
                    "role": "user",
                    "content": [
                        {"type": "image_url", "image_url": {"url": image_data_url(image)}},
                        {"type": "text", "text": prompt},
                    ],
                }
            ],
            model=settings.llm_vl_model,
            temperature=0.0,
            max_tokens=2000,
        )
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"多模态转录失败：{exc}")
        return ""


def _describe_scanned_page(page, page_no: int, document_id: int) -> TextBlock | None:
    """无文字层页面：整页渲染 → 存图 → 多模态转录，产出可检索的文本块。

    这样「被改成图片的表格/流程图」也能进向量库，并保留整页图给模型复核。
    """
    from app.services.media import save_image

    try:
        image = page.to_image(resolution=settings.media_crop_resolution).original
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"整页渲染失败：page={page_no} - {exc}")
        return None

    url = save_image(image, f"{document_id}/p{page_no}_page.png")
    content = _transcribe_image(image, hint=f"文档第 {page_no} 页")
    if not content:
        # 转录失败也要留图：至少前端能看到原页，块内容退回一句说明
        content = f"（第 {page_no} 页为图像，未能识别文字）"

    marked = f"【第 {page_no} 页为图像，以下内容由多模态模型识别】\n{content.strip()}"
    logger.info(f"扫描页转录完成：page={page_no} 字符={len(content)}")
    return TextBlock(
        text=marked,
        page=page_no,
        meta={"source": "vl_page"},
        images=[{"caption": f"第 {page_no} 页（整页图像）", "url": url, "kind": "page"}],
    )


def _page_to_block(page, page_no: int, document_id: int | None) -> TextBlock | None:
    """单页解析：表格转 Markdown、流程图/表格区域裁图。"""
    lines = page.extract_text_lines() or []

    # 1. 表格：抽 markdown，并把被表格覆盖的文本行从正文里剔除，避免重复
    tables = page.find_tables() or []
    table_blocks: list[dict] = []
    for index, table in enumerate(tables, start=1):
        try:
            markdown = _table_to_markdown(table.extract())
        except Exception as exc:  # noqa: BLE001
            logger.warning(f"表格解析失败：page={page_no} - {exc}")
            continue
        if not markdown:
            continue
        caption = markdown.splitlines()[0].strip("| ").replace(" | ", "／")
        table_blocks.append(
            {
                "bbox": tuple(table.bbox),
                "markdown": markdown,
                "top": table.bbox[1],
                "caption": f"表格（第{page_no}页）：{caption[:40]}",
                "index": index,
            }
        )

    def in_table(line: dict) -> bool:
        box = (line["x0"], line["top"], line["x1"], line["bottom"])
        return any(_overlap_ratio(box, item["bbox"]) > 0.5 for item in table_blocks)

    text_parts: list[str] = []
    pending = sorted(table_blocks, key=lambda item: item["top"])
    for line in lines:
        while pending and pending[0]["top"] <= line["top"]:
            text_parts.append(pending.pop(0)["markdown"])
        if in_table(line):
            continue
        stripped = line["text"].strip()
        if stripped:
            text_parts.append(stripped)
    for item in pending:
        text_parts.append(item["markdown"])

    text = "\n".join(text_parts).strip()

    # 0. 整页没有文字层（扫描件 / 被栅格化）：没有文本可抽，交给多模态模型转录
    if document_id and len(text) < _SCAN_TEXT_MIN and getattr(page, "images", None):
        scanned = _describe_scanned_page(page, page_no, document_id)
        if scanned is not None:
            return scanned

    images: list[dict] = []
    if not document_id:
        return TextBlock(text=text, page=page_no) if text else None

    # 2. 表格裁图（模型看表格截图比看压平文本更准）
    for item in table_blocks:
        bbox = (item["bbox"][0] - 4, item["bbox"][1] - 4, item["bbox"][2] + 4, item["bbox"][3] + 4)
        url = _crop_and_save(page, bbox, f"{document_id}/p{page_no}_t{item['index']}.png")
        if url:
            images.append({"caption": item["caption"], "url": url, "kind": "table"})

    # 3. 流程图/示意图：以「图 N」图注所在行为锚点，向上聚类矢量对象
    caption_lines = [
        line for line in lines if _CAPTION_LINE.match(line["text"].strip())
    ]
    for index, line in enumerate(caption_lines, start=1):
        caption = {
            "top": line["top"],
            "bottom": line["bottom"],
            "text": line["text"].strip(),
        }
        bbox = _figure_bbox(page, caption)
        if not bbox:
            continue
        if any(_overlap_ratio(bbox, item["bbox"]) > 0.5 for item in table_blocks):
            continue
        url = _crop_and_save(page, bbox, f"{document_id}/p{page_no}_f{index}.png")
        if url:
            images.append(
                {"caption": caption["text"] or f"图 {index}", "url": url, "kind": "figure"}
            )

    # 4. 位图：没有图注也裁出来，用页码兜底命名
    for index, raster in enumerate(getattr(page, "images", []) or [], start=1):
        bbox = (raster["x0"] - 4, raster["top"] - 4, raster["x1"] + 4, raster["bottom"] + 4)
        if bbox[2] - bbox[0] < 80 or bbox[3] - bbox[1] < 60:
            continue
        if any(_overlap_ratio(bbox, item["bbox"]) > 0.5 for item in table_blocks):
            continue
        url = _crop_and_save(page, bbox, f"{document_id}/p{page_no}_i{index}.png")
        if url:
            images.append({"caption": f"图片（第{page_no}页）", "url": url, "kind": "image"})

    if not text and not images:
        return None
    return TextBlock(text=text, page=page_no, images=images)


def _parse_pdf(path: Path, document_id: int | None = None) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                block = _page_to_block(page, index, document_id)
                if block is not None:
                    blocks.append(block)
    except Exception as exc:  # noqa: BLE001
        raise BizError(f"PDF 解析失败：{exc}") from exc

    if not blocks:
        # 可能是扫描件，尝试用多模态模型识别
        blocks = _ocr_pdf_pages(path, document_id)
    if not blocks:
        raise BizError("PDF 中未提取到任何文字，可能是纯图片扫描件且 OCR 失败")
    return blocks


def _ocr_pdf_pages(path: Path, document_id: int | None = None) -> list[TextBlock]:
    """整本扫描件的兜底：逐页渲染后交给多模态模型转录（不依赖 PyMuPDF）。"""
    blocks: list[TextBlock] = []
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                if index > 10:  # 扫描件最多识别前 10 页，避免耗时过长
                    break
                if document_id:
                    block = _describe_scanned_page(page, index, document_id)
                    if block is not None:
                        blocks.append(block)
                        continue
                try:
                    image = page.to_image(resolution=settings.media_crop_resolution).original
                except Exception as exc:  # noqa: BLE001
                    logger.warning(f"扫描页渲染失败：page={index} - {exc}")
                    continue
                text = _transcribe_image(image, hint=f"文档第 {index} 页")
                if text:
                    blocks.append(TextBlock(text=text, page=index, meta={"source": "vl_page"}))
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"扫描件识别失败：{exc}")
    return blocks


def _parse_pptx(path: Path) -> list[TextBlock]:
    try:
        from pptx import Presentation

        presentation = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001
        raise BizError(f"PPT 解析失败：{exc}") from exc

    blocks: list[TextBlock] = []
    for index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        parts.append(line)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"备注：{notes}")
        if parts:
            blocks.append(TextBlock(text="\n".join(parts), page=index))
    if not blocks:
        raise BizError("PPT 中未提取到任何文字内容")
    return blocks


def _parse_docx(path: Path) -> list[TextBlock]:
    try:
        import docx

        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise BizError(f"Word 解析失败：{exc}") from exc

    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                paragraphs.append(line)
    if not paragraphs:
        raise BizError("Word 文档中未提取到任何文字内容")
    return [TextBlock(text="\n".join(paragraphs))]


def _parse_text(path: Path) -> list[TextBlock]:
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        content = path.read_text(encoding="gbk", errors="ignore")
    content = content.strip()
    if not content:
        raise BizError("文档内容为空")
    return _split_markdown(content)


_MD_H2 = re.compile(r"^##\s+")


def _split_markdown(content: str) -> list[TextBlock]:
    """把 Markdown 按二级标题拆成多个解析单元（每个单元 = 一个父块）。

    原来的实现让整篇文档成为唯一父块，父块粒度等于文档粒度，
    检索「命中一段」和「命中整篇」没有区别，评估指标会被系统性拉饱和。
    这里改为按 `##` 章节切父块：子块仍在章节内部细切，父块回填的是该章节全文。
    """
    lines = content.splitlines()
    title = ""
    body_start = 0
    for index, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("# ") and not stripped.startswith("## "):
            title = stripped[2:].strip()
            body_start = index + 1
            break

    body = "\n".join(lines[body_start:]).strip()
    if not body:
        body = content

    sections: list[str] = []
    current: list[str] = []
    for line in body.splitlines():
        if _MD_H2.match(line):
            if current:
                sections.append("\n".join(current).strip())
            current = [line]
        else:
            current.append(line)
    if current:
        sections.append("\n".join(current).strip())
    sections = [item for item in sections if item]

    # 首个 `##` 之前的内容（文档说明、免责声明等）挂到第一节，避免产生只有一句话的碎父块
    if sections and not _MD_H2.match(sections[0]):
        head = sections.pop(0)
        if sections:
            sections[0] = f"{head}\n\n{sections[0]}"
        else:
            sections = [head]

    if len(sections) <= 1:
        return [TextBlock(text=body, meta={"title": title} if title else {})]

    return [
        TextBlock(text=section, meta={"title": title, "section": order})
        for order, section in enumerate(sections, start=1)
    ]


def _parse_image(path: Path) -> list[TextBlock]:
    text = _ocr_image_bytes(path.read_bytes())
    if not text:
        raise BizError("图片 OCR 未识别到文字内容")
    return [TextBlock(text=text)]


def _ocr_image_bytes(image_bytes: bytes) -> str:
    """调用通义千问多模态模型做 OCR，返回识别出的文字。"""
    from app.core.config import settings
    from app.services.llm import get_client, thinking_kwargs

    encoded = base64.b64encode(image_bytes).decode("ascii")
    client = get_client()
    errors: list[str] = []
    # 专用 OCR 模型优先；网关没部署时退回通用视觉模型，避免整篇文档解析失败
    for model in dict.fromkeys([settings.vision_model, settings.llm_vl_model]):
        try:
            response = client.chat.completions.create(
                model=model,
                messages=[
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "image_url",
                                "image_url": {"url": f"data:image/png;base64,{encoded}"},
                            },
                            {
                                "type": "text",
                                "text": "请提取图片中的全部文字，保持原有段落顺序，只输出文字内容，不要添加解释。",
                            },
                        ],
                    }
                ],
                timeout=settings.llm_timeout,
                **thinking_kwargs(model),
            )
            return (response.choices[0].message.content or "").strip()
        except Exception as exc:  # noqa: BLE001
            errors.append(f"{model}: {exc}")
    raise BizError(f"图片文字识别失败：{'; '.join(errors)[:300]}")
